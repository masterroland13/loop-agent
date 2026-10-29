# -*- coding: utf-8 -*-
"""Human Files panel APIs: list / read (bare text) / write — sandbox only; not propose_edit."""

from __future__ import annotations

from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from agent_window.config import settings
from agent_window.main import app
from agent_window.services.workspace import (
    WorkspaceError,
    clear_workspace_root,
    copy_entry,
    delete_entry,
    list_directory,
    mkdir_directory,
    read_file,
    read_docx_file,
    read_editor_file,
    read_text_file,
    rename_entry,
    reveal_in_explorer,
    set_workspace_root,
    tree_stats,
    write_docx_file,
    write_editor_file,
    write_text_file,
)


@pytest.fixture
def ws(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    data = tmp_path / "data"
    data.mkdir()
    monkeypatch.setattr(settings, "app_data_dir", str(data))
    monkeypatch.setattr("agent_window.models.settings.app_data_dir", str(data))
    monkeypatch.setattr("agent_window.services.storage.settings.app_data_dir", str(data))
    root = tmp_path / "ws"
    root.mkdir()
    set_workspace_root(str(root))
    yield root
    clear_workspace_root()


@pytest.fixture
def client(ws: Path) -> TestClient:
    with TestClient(app) as c:
        yield c


def test_read_text_file_no_line_numbers(ws: Path) -> None:
    (ws / "a.txt").write_text("hello\nworld\n", encoding="utf-8")
    bare = read_text_file("a.txt")
    assert bare["content"] == "hello\nworld\n"
    assert "1|" not in bare["content"]
    assert bare["path"] == "a.txt"
    assert bare["mtime_ns"] > 0
    # Agent read_file still numbered — must not regress.
    agent = read_file("a.txt", offset=1, limit=10)
    assert "1|hello" in agent["content"]
    assert agent["content"] != bare["content"]


def test_write_text_file_roundtrip(ws: Path) -> None:
    write_text_file("b.txt", "one\n")
    assert (ws / "b.txt").read_text(encoding="utf-8") == "one\n"
    again = write_text_file("b.txt", "two\n", if_match_mtime_ns=read_text_file("b.txt")["mtime_ns"])
    assert again["size"] == 4
    assert (ws / "b.txt").read_text(encoding="utf-8") == "two\n"


def test_write_mtime_conflict(ws: Path) -> None:
    import os

    write_text_file("c.txt", "v1\n")
    stale = read_text_file("c.txt")["mtime_ns"]
    (ws / "c.txt").write_text("external\n", encoding="utf-8")
    # Coarse FS mtime (e.g. 1s) may not advance on fast rewrite — force a newer stamp.
    newer = stale + 2_000_000_000
    os.utime(ws / "c.txt", ns=(newer, newer))
    with pytest.raises(WorkspaceError) as ei:
        write_text_file("c.txt", "mine\n", if_match_mtime_ns=stale)
    assert ei.value.code == "CONFLICT"
    assert ei.value.code != "OK"


def test_write_mtime_drift_unchanged_base_content_ok(ws: Path) -> None:
    import os

    write_text_file("d.txt", "base\n")
    stale = read_text_file("d.txt")["mtime_ns"]
    newer = stale + 2_000_000_000
    os.utime(ws / "d.txt", ns=(newer, newer))
    result = write_text_file(
        "d.txt",
        "edited\n",
        if_match_mtime_ns=stale,
        if_match_base_content="base\n",
    )
    assert result["size"] == 7
    assert (ws / "d.txt").read_text(encoding="utf-8") == "edited\n"


def test_write_mtime_drift_changed_content_conflict(ws: Path) -> None:
    import os

    write_text_file("e.txt", "base\n")
    stale = read_text_file("e.txt")["mtime_ns"]
    (ws / "e.txt").write_text("external\n", encoding="utf-8")
    newer = stale + 2_000_000_000
    os.utime(ws / "e.txt", ns=(newer, newer))
    with pytest.raises(WorkspaceError) as ei:
        write_text_file(
            "e.txt",
            "mine\n",
            if_match_mtime_ns=stale,
            if_match_base_content="base\n",
        )
    assert ei.value.code == "CONFLICT"


def test_path_escape_rejected(ws: Path) -> None:
    with pytest.raises(WorkspaceError) as ei:
        read_text_file("../outside.txt")
    assert ei.value.code == "PATH_ESCAPE"


def test_no_workspace_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    data = tmp_path / "data"
    data.mkdir()
    monkeypatch.setattr(settings, "app_data_dir", str(data))
    clear_workspace_root()
    with pytest.raises(WorkspaceError) as ei:
        list_directory(".")
    assert ei.value.code == "NO_WORKSPACE"


def test_binary_rejected(ws: Path) -> None:
    (ws / "bin.dat").write_bytes(b"abc\x00def")
    with pytest.raises(WorkspaceError) as ei:
        read_text_file("bin.dat")
    assert ei.value.code == "BINARY"


def test_http_list_read_write(client: TestClient, ws: Path) -> None:
    (ws / "note.md").write_text("# hi\n", encoding="utf-8")
    listed = client.get("/api/workspace/files/list", params={"path": "."})
    assert listed.status_code == 200
    names = {e["name"] for e in listed.json()["entries"]}
    assert "note.md" in names
    assert "note.md" not in {""}

    read = client.get("/api/workspace/files/read", params={"path": "note.md"})
    assert read.status_code == 200
    body = read.json()
    assert body["content"] == "# hi\n"
    assert "1|" not in body["content"]
    mtime = body["mtime_ns"]
    assert isinstance(mtime, str)
    assert mtime.isdigit()

    written = client.put(
        "/api/workspace/files/write",
        json={"path": "note.md", "content": "# bye\n", "if_match_mtime_ns": mtime},
    )
    assert written.status_code == 200
    assert (ws / "note.md").read_text(encoding="utf-8") == "# bye\n"

    conflict = client.put(
        "/api/workspace/files/write",
        json={"path": "note.md", "content": "x", "if_match_mtime_ns": mtime},
    )
    assert conflict.status_code == 409
    assert conflict.json()["detail"]["code"] == "CONFLICT"


def test_http_escape(client: TestClient, ws: Path) -> None:
    r = client.get("/api/workspace/files/read", params={"path": "../x"})
    assert r.status_code == 400
    assert r.json()["detail"]["code"] == "PATH_ESCAPE"


def test_mkdir_directory(ws: Path) -> None:
    created = mkdir_directory("docs")
    assert created["path"] == "docs"
    assert (ws / "docs").is_dir()
    with pytest.raises(WorkspaceError) as dup:
        mkdir_directory("docs")
    assert dup.value.code == "EXISTS"


def test_rename_file(ws: Path) -> None:
    (ws / "old.txt").write_text("x", encoding="utf-8")
    out = rename_entry("old.txt", "new.txt")
    assert out["path"] == "new.txt"
    assert not (ws / "old.txt").exists()
    assert (ws / "new.txt").read_text(encoding="utf-8") == "x"


def test_delete_file_and_empty_dir(ws: Path) -> None:
    (ws / "gone.txt").write_text("bye", encoding="utf-8")
    assert delete_entry("gone.txt")["path"] == "gone.txt"
    assert not (ws / "gone.txt").exists()

    (ws / "empty").mkdir()
    assert delete_entry("empty")["path"] == "empty"
    assert not (ws / "empty").exists()


def test_delete_nonempty_dir_rejected(ws: Path) -> None:
    d = ws / "full"
    d.mkdir()
    (d / "a.txt").write_text("a", encoding="utf-8")
    with pytest.raises(WorkspaceError) as ei:
        delete_entry("full")
    assert ei.value.code == "NOT_EMPTY"
    assert (d / "a.txt").exists()


def test_tree_stats(ws: Path) -> None:
    (ws / "pkg").mkdir()
    (ws / "pkg" / "sub").mkdir()
    (ws / "pkg" / "a.txt").write_text("a", encoding="utf-8")
    (ws / "pkg" / "sub" / "b.txt").write_text("b", encoding="utf-8")
    stats = tree_stats("pkg")
    assert stats["file_count"] == 2
    assert stats["dir_count"] == 1


def test_copy_file_and_directory(ws: Path) -> None:
    (ws / "src.txt").write_text("hello", encoding="utf-8")
    copied = copy_entry("src.txt", "dst.txt")
    assert copied["path"] == "dst.txt"
    assert (ws / "dst.txt").read_text(encoding="utf-8") == "hello"
    assert (ws / "src.txt").exists()

    (ws / "tree").mkdir()
    (ws / "tree" / "leaf.txt").write_text("x", encoding="utf-8")
    tree_copy = copy_entry("tree", "tree-copy")
    assert tree_copy["path"] == "tree-copy"
    assert (ws / "tree-copy" / "leaf.txt").read_text(encoding="utf-8") == "x"


def test_copy_into_descendant_rejected(ws: Path) -> None:
    (ws / "tree").mkdir()
    with pytest.raises(WorkspaceError) as ei:
        copy_entry("tree", "tree/child")
    assert ei.value.code == "INVALID"


def test_delete_recursive_directory(ws: Path) -> None:
    (ws / "tree").mkdir()
    (ws / "tree" / "nested").mkdir()
    (ws / "tree" / "nested" / "x.txt").write_text("x", encoding="utf-8")
    out = delete_entry("tree", recursive=True)
    assert out["recursive"] is True
    assert not (ws / "tree").exists()


def test_reveal_in_explorer_xdg(ws: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    target = ws / "show.me"
    target.write_text("x", encoding="utf-8")
    calls: list[list[str]] = []

    def fake_popen(args, **kwargs):
        calls.append(list(args))
        return None

    monkeypatch.setattr("agent_window.services.workspace.subprocess.Popen", fake_popen)
    monkeypatch.setattr("agent_window.services.workspace.platform.system", lambda: "Linux")
    monkeypatch.setattr("agent_window.services.workspace._in_wsl", lambda: False)
    monkeypatch.setattr(
        "agent_window.services.workspace.shutil.which",
        lambda cmd: "/usr/bin/xdg-open" if cmd == "xdg-open" else None,
    )
    data = reveal_in_explorer("show.me")
    assert data["path"] == "show.me"
    assert calls and calls[0][0] == "/usr/bin/xdg-open"


def test_reveal_in_explorer_wsl(ws: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    target = ws / "show.me"
    target.write_text("x", encoding="utf-8")
    calls: list[list[str]] = []

    monkeypatch.setattr(
        "agent_window.services.workspace.subprocess.Popen",
        lambda args, **kwargs: calls.append(list(args)),
    )
    monkeypatch.setattr("agent_window.services.workspace.platform.system", lambda: "Linux")
    monkeypatch.setattr("agent_window.services.workspace._in_wsl", lambda: True)
    monkeypatch.setattr(
        "agent_window.services.workspace._to_windows_path",
        lambda p: "C:\\\\wsl$\\\\show.me",
    )
    monkeypatch.setattr(
        "agent_window.services.workspace._explorer_exe",
        lambda: "explorer.exe",
    )
    data = reveal_in_explorer("show.me")
    assert data["path"] == "show.me"
    assert calls[0][:2] == ["explorer.exe", "/select,"]


def test_http_mkdir_rename_delete(client: TestClient, ws: Path) -> None:
    mk = client.post("/api/workspace/files/mkdir", json={"path": "pkg"})
    assert mk.status_code == 200
    assert mk.json()["path"] == "pkg"
    assert (ws / "pkg").is_dir()

    (ws / "pkg" / "a.txt").write_text("1", encoding="utf-8")
    rn = client.post(
        "/api/workspace/files/rename",
        json={"from": "pkg/a.txt", "to": "pkg/b.txt"},
    )
    assert rn.status_code == 200
    renamed = rn.json()
    assert renamed["path"] == "pkg/b.txt"
    assert renamed["from"] == "pkg/a.txt"
    assert "from_path" not in renamed
    assert not (ws / "pkg" / "a.txt").exists()

    de = client.delete("/api/workspace/files/delete", params={"path": "pkg/b.txt"})
    assert de.status_code == 200
    assert not (ws / "pkg" / "b.txt").exists()

    de_dir = client.delete("/api/workspace/files/delete", params={"path": "pkg"})
    assert de_dir.status_code == 200
    assert not (ws / "pkg").exists()


def test_http_copy_stats_reveal_recursive(client: TestClient, ws: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    (ws / "box").mkdir()
    (ws / "box" / "one.txt").write_text("1", encoding="utf-8")

    stats = client.get("/api/workspace/files/stats", params={"path": "box"})
    assert stats.status_code == 200
    body = stats.json()
    assert body["file_count"] == 1
    assert body["dir_count"] == 0

    copied = client.post(
        "/api/workspace/files/copy",
        json={"from": "box/one.txt", "to": "box/two.txt"},
    )
    assert copied.status_code == 200
    copied_body = copied.json()
    assert copied_body["path"] == "box/two.txt"
    assert copied_body["from"] == "box/one.txt"
    assert "from_path" not in copied_body
    assert (ws / "box" / "two.txt").exists()

    calls: list[list[str]] = []
    monkeypatch.setattr(
        "agent_window.services.workspace.subprocess.Popen",
        lambda args, **kwargs: calls.append(list(args)),
    )
    monkeypatch.setattr("agent_window.services.workspace.platform.system", lambda: "Linux")
    revealed = client.post("/api/workspace/files/reveal", json={"path": "box/one.txt"})
    assert revealed.status_code == 200
    assert calls

    (ws / "tree").mkdir()
    (ws / "tree" / "nested").mkdir()
    (ws / "tree" / "nested" / "x.txt").write_text("x", encoding="utf-8")
    deleted = client.delete("/api/workspace/files/delete", params={"path": "tree", "recursive": True})
    assert deleted.status_code == 200
    assert deleted.json()["recursive"] is True
    assert not (ws / "tree").exists()


def test_read_write_docx_html(ws: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello Word")
    doc.save(ws / "letter.docx")

    data = read_docx_file("letter.docx")
    assert data["encoding"] == "docx-html"
    assert "Hello" in data["content"]

    updated = write_docx_file(
        "letter.docx",
        "<p>Hello Word updated</p>",
        if_match_mtime_ns=data["mtime_ns"],
        if_match_base_content=data["content"],
    )
    assert updated["size"] > 0
    again = read_docx_file("letter.docx")
    assert "updated" in again["content"]


def test_http_read_write_docx(client: TestClient, ws: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("HTTP docx")
    doc.save(ws / "http.docx")

    read = client.get("/api/workspace/files/read", params={"path": "http.docx"})
    assert read.status_code == 200
    body = read.json()
    assert body["encoding"] == "docx-html"
    assert "HTTP" in body["content"]
    mtime = body["mtime_ns"]

    written = client.put(
        "/api/workspace/files/write",
        json={
            "path": "http.docx",
            "content": "<p>HTTP docx saved</p>",
            "if_match_mtime_ns": mtime,
            "if_match_base_content": body["content"],
            "encoding": "docx-html",
        },
    )
    assert written.status_code == 200
    reread = read_editor_file("http.docx")
    assert "saved" in reread["content"]
